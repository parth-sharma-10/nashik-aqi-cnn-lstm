"""Build the DOCX report from outputs/backtest_2025/ and outputs/forecast_2026/.

Every number in the text is read from those files. Qualitative statements that depend on
the results are guarded by asserts, so a different run fails loudly instead of
misreporting. Run from the project root after both modes:  python scripts/build_report.py
"""
import json

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor

REPORT = "outputs/Nashik_AQ_Forecast_2026_Report.docx"
POLL = ["PM2.5", "PM10", "NO2", "SO2", "CO", "O3", "NH3"]
STATIONS = ["MIDC Ambad", "Gangapur Road", "Pandav Nagari", "Hirwadi"]
MODES = {"backtest_2025": "2025 backtest", "forecast_2026": "2026 forecast"}


def load(mode):
    out = f"outputs/{mode}"
    r = {"out": out, "meta": json.load(open(f"{out}/run_metadata.json"))}
    for key, name in [("rec", "recursive_station_pollutant_metrics"),
                      ("rec_all", "recursive_overall_pollutant_metrics"),
                      ("one", "one_step_station_pollutant_metrics"),
                      ("one_all", "one_step_overall_pollutant_metrics"),
                      ("monthly", "recursive_monthly_mean_metrics"),
                      ("aqi", "aqi_summary")]:
        r[key] = pd.read_csv(f"{out}/metrics/{name}.csv")
    r["sel"] = pd.read_csv(f"{out}/metrics/model_selection_validation.csv", index_col="station")
    r["splits"] = pd.read_csv(f"{out}/tables/splits.csv")
    r["cov"] = pd.read_csv(f"{out}/tables/horizon_coverage.csv")
    return r


B, F = load("backtest_2025"), load("forecast_2026")
insp = pd.read_csv("outputs/forecast_2026/tables/dataset_inspection.csv")
clog = pd.read_csv("outputs/forecast_2026/tables/cleaning_log.csv", index_col=0)
imp = pd.read_csv("outputs/forecast_2026/tables/imputation_counts.csv")
miss = pd.read_csv("outputs/forecast_2026/tables/missing_pct_before_imputation.csv", index_col=0)
V = F["meta"]["versions"]


def pooled(r, kind, col):
    t = r[f"{kind}_all"].set_index("pollutant").loc[POLL]
    return t[col]


def skill_range(r, kind, ref):
    s = pooled(r, kind, f"RMSE skill vs {ref}")
    return (f"{100 * s.min():.1f}% ({s.idxmin()}) to {100 * s.max():.1f}% ({s.idxmax()})",
            s)


B_rec_skill_txt, B_rec_skill = skill_range(B, "rec", "climatology")
F_rec_skill_txt, F_rec_skill = skill_range(F, "rec", "climatology")
B_one_skill_txt, B_one_skill = skill_range(B, "one", "persistence")
F_one_skill_txt, F_one_skill = skill_range(F, "one", "persistence")


def wins(r, kind, ref):
    """(series where the CNN-LSTM beats the baseline, scored unflagged series, median skill)."""
    t = r[kind]
    s = t[(t.n_valid > 0) & ~t.quality_flag][f"RMSE skill vs {ref}"]
    return int((s > 0).sum()), len(s), float(s.median())


B_one_wins, B_one_n, B_one_median = wins(B, "one", "persistence")
F_one_wins, F_one_n, F_one_median = wins(F, "one", "persistence")
B_rec_wins, B_rec_n, _ = wins(B, "rec", "climatology")
F_rec_wins, F_rec_n, _ = wins(F, "rec", "climatology")
F_measured = F["cov"].set_index("station")
fc26 = pd.read_csv(f"{F['out']}/predictions/recursive_forecast.csv",
                   usecols=["station", "pollutant", "month", "cnn_lstm_pred", "climatology_pred"])
pm_monthly = fc26[fc26.pollutant == "PM2.5"].groupby(["station", "month"])[
    ["cnn_lstm_pred", "climatology_pred"]].mean().unstack("month").reindex(STATIONS)
pm_ratio = pm_monthly["cnn_lstm_pred"] / pm_monthly["climatology_pred"]


def count_txt(s, better):
    n = int((s > 0).sum()) if better else int((s < 0).sum())
    return "every pollutant" if n == len(s) else f"{n} of {len(s)} pollutants"


B_rec_worse_txt = count_txt(B_rec_skill, better=False)
F_rec_worse_txt = count_txt(F_rec_skill, better=False)


def worst_series(r):
    t = r["rec"][(r["rec"].n_valid > 0) & ~r["rec"].quality_flag]
    w = t.loc[t["RMSE skill vs climatology"].idxmin()]
    return (f"{w.station} {w.pollutant} (RMSE {w['CNN-LSTM RMSE']:.1f} against "
            f"{w['Climatology RMSE']:.1f} for climatology)")


# Facts the prose below states. If a re-run changes them, rewrite the prose.
assert F_measured.loc["Pandav Nagari", "rows_in_file"] == 0
assert (F_measured.drop("Pandav Nagari").rows_in_file == 5088).all()
assert all("Jul" in F_measured.loc[s, "months_without_rows"] for s in STATIONS)
assert (B["cov"].rows_in_file == 8760).all()
assert (F["rec"][F["rec"].station == "Pandav Nagari"].n_valid == 0).all()
# The recursive forecast is worse than climatology for most pollutants, pooled, in both years;
# the one-step predictions beat persistence for most pollutants and most series in both years.
assert (B_rec_skill < 0).sum() > len(POLL) / 2 and (F_rec_skill < 0).sum() > len(POLL) / 2
assert B_rec_wins < B_rec_n / 2 and F_rec_wins < F_rec_n / 2
assert (B_one_skill > 0).sum() > len(POLL) / 2 and (F_one_skill > 0).sum() > len(POLL) / 2
assert B_one_wins > B_one_n / 2 and F_one_wins > F_one_n / 2
assert ((pm_ratio > 2) | (pm_ratio < 0.5)).any().any()  # drift stated in Section 12.1

doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(11)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(1)
FIG_N = [0]
TAB_N = [0]


def h(text, level=1):
    doc.add_heading(text, level=level)


def para(text, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text)
    r.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def equation(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name, r.italic = "Cambria Math", True


def caption(prefix, counter, text):
    counter[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f"{prefix} {counter[0]}. {text}")
    r.italic, r.font.size = True, Pt(9)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return counter[0]


def figure(mode_out, path, text, width=6.3):
    doc.add_picture(f"{mode_out}/figures/{path}", width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return caption("Figure", FIG_N, text)


def cell_text(v):
    if isinstance(v, str):
        return v
    if isinstance(v, float):
        return "-" if pd.isna(v) else f"{v:,.4g}"
    return str(v)


def table(df, text, font=8, widths=None):
    n = caption("Table", TAB_N, text)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style, t.alignment = "Light Grid Accent 1", WD_TABLE_ALIGNMENT.CENTER
    for cell, col in zip(t.rows[0].cells, df.columns):
        cell.text = str(col)
    for _, row in df.iterrows():
        for cell, v in zip(t.add_row().cells, row):
            cell.text = cell_text(v)
    for row in t.rows:
        for k, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[k])
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
    doc.add_paragraph()
    return n


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def flag(df):
    return [f"{p} *" if q else p for p, q in zip(df.pollutant, df.quality_flag)]


def front_matter():
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Hourly Air-Quality Forecasting for Nashik Monitoring Stations with "
                  "KNN Imputation and a CNN-LSTM Network: 2025 Backtest and Partial 2026 "
                  "Evaluation")
    r.bold, r.font.size = True, Pt(20)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Academic Project Report").font.size = Pt(13)
    for line in ["", "[Student name]  ·  [Roll number]", "[Department, Institution]",
                 "[Course / guide]", "[Submission date]"]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_break()

    h("Abstract")
    para(
        "Hourly concentrations of seven pollutants (PM2.5, PM10, NO2, SO2, CO, O3 and NH3) "
        "from four continuous monitoring stations in Nashik, India (MIDC Ambad, Gangapur Road, "
        "Pandav Nagari and Hirwadi) were used to train one convolutional-LSTM (CNN-LSTM) "
        "network per station that predicts the next hour from the previous 24. Missing inputs "
        "were filled with a K-nearest-neighbours imputer fitted on training data only. The "
        "pipeline runs in two separate modes. The 2025 backtest uses data up to "
        "31 December 2024 (network fitted to 30 June 2024, validated on July-December 2024) "
        "and is scored on 2025, which is measured at all stations. The 2026 run uses data "
        "up to 31 December 2025 in the same way and forecasts all "
        "8,760 hours of 2026 at all four stations; measurements exist only for January-June "
        "and August 2026 at three stations and not at all at Pandav Nagari. Two kinds of "
        "prediction are scored separately on measured hours: a recursive forecast of the "
        "whole year from the forecast origin, which uses no later measurement, and a rolling "
        "one-step evaluation, in which the frozen model predicts each measured hour from the "
        "measured previous 24 hours. Faulty sensor readings (stuck values, CO zeros and "
        "near-constant 'flatline' days) are removed before training and are never scored. "
        "Pooled over stations, with quality-flagged sensor series excluded, the recursive "
        "forecast changed RMSE relative to a month-by-hour climatology by "
        f"{B_rec_skill_txt} in 2025 and {F_rec_skill_txt} in 2026 (negative = worse than "
        "climatology): fed its own predictions for a year, the network is a worse year-ahead "
        f"estimate than climatology for {B_rec_worse_txt} in 2025 and {F_rec_worse_txt} in "
        "2026. The one-step predictions changed RMSE relative to persistence by "
        f"{B_one_skill_txt} in 2025 and {F_one_skill_txt} in 2026: a small, mostly "
        "consistent gain when fresh measurements are available every hour. Daily CPCB "
        "AQI values were derived from the predicted concentrations; they are model estimates, "
        "not official AQI bulletins.")
    h("Keywords")
    para("air-quality forecasting; CNN-LSTM; KNN imputation; recursive forecasting; backtest; "
         "CPCB AQI; Nashik; chronological validation")
    page_break()
    key_findings()
    page_break()


def key_findings():
    h("Key Findings in Plain Language")
    para("This page states the conclusions first. The rest of the report gives the evidence.")
    bullets([
        "What the model does. For each station, a neural network reads the last 24 hours "
        "of seven pollutant measurements and predicts the next hour.",
        "Next-hour prediction works, modestly. The simplest possible guess is persistence: "
        "'next hour = the latest measured value'. One hour ahead, the CNN-LSTM is more "
        f"accurate than persistence for {B_one_wins} of {B_one_n} station-pollutant series "
        f"in 2025 and {F_one_wins} of {F_one_n} in 2026. The typical (median) error "
        f"reduction is {100 * B_one_median:.1f}% in 2025 and {100 * F_one_median:.1f}% in "
        "2026. That is a real but small improvement.",
        "A high R² is not proof of a good model. Hourly pollution changes slowly, so "
        "persistence already reaches a high R². Every R² in this report is shown next to a "
        "baseline scored on exactly the same hours.",
        "A full-year forecast does not work with this model. Feeding the model its own "
        "predictions for a whole year lets errors pile up. Against climatology (the "
        "average for that month and hour of day), the recursive CNN-LSTM is better for "
        f"only {B_rec_wins} of {B_rec_n} series in 2025 and {F_rec_wins} of {F_rec_n} in "
        "2026. For months ahead, climatology is the better estimate.",
        "Faulty sensor readings were removed before training and scoring. Examples are "
        "Gangapur Road NO2 stuck near 3.85 µg/m³ from late August 2025, Gangapur Road NH3 "
        "stuck near 12 µg/m³ in 2026, and CO reading 0.00 mg/m³ for months in 2026. Every "
        "removed hour is counted and labelled; pollution peaks are never removed.",
        "2026 is only a partial test: January-June and August at three stations. July is "
        "missing from all files, September-December have not happened in the data, and "
        "Pandav Nagari has no 2026 data at all.",
    ])
    h("Glossary", 2)
    bullets([
        "Station-pollutant series: the hourly record of one pollutant at one station "
        "(4 stations x 7 pollutants = 28 series).",
        "Persistence baseline: predict that the next hour equals the latest measured value.",
        "Climatology baseline: predict the historical average for that calendar month and "
        "hour of day.",
        "RMSE / MAE: typical size of the prediction error, in µg/m³ (mg/m³ for CO). "
        "Lower is better.",
        "Skill: 1 - RMSE(model) / RMSE(baseline). Positive means the model beats the "
        "baseline; 0.05 means 5% lower error.",
        "Rolling one-step prediction: each hour is predicted from the real measurements of "
        "the previous 24 hours.",
        "Recursive forecast: the model's own predictions are fed back as inputs, so it can "
        "run for a whole year without any future measurement.",
        "Imputation: filling a missing input value with an estimate. Missing target values "
        "are never filled; they are simply not scored.",
    ])


def intro_sections():
    h("1. Introduction")
    para("Short-term air-quality prediction supports public advisories and operational "
         "decisions such as traffic or industrial-emission controls. Continuous ambient "
         "monitoring stations report hourly concentrations, but their records contain gaps, "
         "instrument faults and changes in data availability, so any forecasting pipeline must "
         "handle missing data without using information from after the forecast origin.")
    para("Hybrid convolutional-recurrent networks are a common choice for such series: "
         "one-dimensional convolutions extract short local patterns, and Long Short-Term "
         "Memory (LSTM) layers model longer dependencies [1], [3], [4]. This project applies "
         "that design to four Nashik stations, produces an hourly forecast of 2026, and "
         "tests the same procedure on 2025, where the full year was measured.")
    h("2. Problem Statement")
    para("Using only the hourly observations of a station available before 1 January 2026, "
         "forecast the concentrations of seven pollutants for every hour of 2026, and measure "
         "the accuracy of the same procedure where the answer is known: on 2025, with data "
         "up to 31 December 2024, and on the hours of 2026 that have since been measured.")
    h("3. Objectives")
    bullets([
        "Document the dataset's schema, coverage and quality problems per station.",
        "Build reproducible cleaning and KNN imputation fitted without temporal leakage.",
        "Train one CNN-LSTM per station on data before each forecast origin.",
        "Export a continuous hourly 2026 forecast for all four stations, including hours "
        "and stations without measurements, without fabricating ground truth.",
        "Score recursive and one-step predictions separately with MSE, RMSE, MAE and R² in "
        "original units, against climatology, same-hour-last-year and persistence baselines.",
        "Derive daily CPCB AQI from the predicted concentrations, clearly separated from "
        "measured AQI.",
    ])


def dataset_section():
    h("4. Dataset Description")
    para("The dataset consists of four CSV files, one per station. Station identity is "
         "given only by the file name. All files share 25 columns: a Timestamp plus 24 "
         "measurements (criteria pollutants, volatile organic compounds and meteorological "
         "variables). The Hirwadi file has an extra Source_File column whose labels do not "
         "match the timestamps, so it was dropped and the timestamps were treated as "
         "authoritative. Hirwadi rows are stored out of order and were sorted. No file "
         "contains duplicate timestamps.")
    d = insp[["station", "rows", "first", "last", "duplicate_timestamps", "sorted_in_file",
              "years_present", "all_pollutants_missing_rows"]].copy()
    d.columns = ["Station", "Rows", "First timestamp", "Last timestamp", "Duplicate times",
                 "Sorted", "Years present", "Rows with no pollutant"]
    table(d, "Raw files as delivered.", font=7)
    c = pd.concat([
        B["cov"].assign(year=2025), F["cov"].assign(year=2026)])[
        ["station", "year", "rows_in_file", "hours_any_pollutant_measured", "first_measured",
         "last_measured", "months_without_rows"]]
    c.columns = ["Station", "Horizon", "Rows in file (of 8,760)", "Hours with a measurement",
                 "First measured", "Last measured", "Months with no rows"]
    table(c, "Coverage of each forecast horizon in the station files.", font=7)
    para("Three files begin with an empty 2017 block and resume in mid-2023; Gangapur Road "
         "has data from 2017 onwards. In 2026, every file that has 2026 rows covers "
         "1 January-30 June and 1-31 August (5,088 of 8,760 hours); July and September-"
         "December are absent. Pandav Nagari's file ends on 31 December 2025. Its 2026 "
         "forecast is produced like the others but cannot be scored.")
    para("Seven pollutants are the prediction targets: the pollutants of the Indian National "
         "AQI present in the data (PM2.5, PM10, NO2, SO2, CO, O3, NH3). The model uses no "
         "weather inputs. The files do contain weather columns, which are measured for "
         "about 93-97% of hours at MIDC Ambad, Pandav Nagari and Hirwadi. At Gangapur Road, "
         "temperature and humidity cover only about 5% of hours. Weather was left out to "
         "keep one consistent input set for all stations.")
    para("In January-May 2026 many values at the three stations with 2026 data are reported "
         "as whole numbers (25-50% of values, against about 1% before). This looks like a "
         "change in reporting precision, not a sensor fault, and these values are kept.")


def preprocessing_section():
    h("5. Data Preprocessing and KNN Imputation")
    h("5.1 Cleaning rules", 2)
    n_ceiling = int(clog["PM2.5_ceiling"].sum() + clog["PM10_ceiling"].sum())
    bullets([
        "Timestamps parsed and sorted; a duplicate timestamp would keep its first row "
        "(none occur).",
        "Rows before a station's first pollutant observation removed (the empty 2017 blocks).",
        "Negative concentrations set to missing (none occur).",
        f"PM2.5 or PM10 of 999 or more set to missing: the analyser ceiling ({n_ceiling} "
        "values). The true value is unknown above the ceiling; one such hour is Diwali "
        "night 2024 at Gangapur Road.",
        "Any value repeated unchanged for 12 or more consecutive hours set to missing as a "
        f"stuck-sensor reading ({int(clog.filter(like='_stuck').sum().sum()):,} values).",
        "CO of exactly 0.00 mg/m³ set to missing. Outdoor CO never falls to zero (the "
        "clean-air background is about 0.1 mg/m³), and the zeros come in long blocks when "
        f"the analyser is off ({int(clog['CO_zero'].sum()):,} further values after the "
        "stuck rule). Zeros of other pollutants are kept.",
        "Flatline days set to missing. A failed sensor can jitter around a fixed level "
        "instead of repeating one value exactly. For each station, pollutant and calendar "
        "day with at least 12 measured hours, the day is removed when the spread of its "
        "values (90th minus 10th percentile) is below 5% of its median "
        f"({int(clog.filter(like='_flatline').sum().sum()):,} values). Real pollution "
        "varies far more within a day, so this rule only removes near-constant days and "
        "can never remove a pollution peak.",
        "No interpolation or resampling. Hours absent from a file are not filled, and no "
        "input window spans them.",
    ])
    q = pd.DataFrame({rule: clog.filter(like=f"_{rule}").sum(axis=1).astype(int)
                      for rule in ("ceiling", "stuck", "zero", "flatline")})
    q.insert(0, "Station", q.index)
    q.columns = ["Station", "PM ceiling", "Stuck 12 h", "CO zero", "Flatline days"]
    table(q, "Hourly values removed by each cleaning rule, whole file (all pollutants).")
    flat = clog.filter(like="_flatline").copy()
    flat.columns = [c.replace("_flatline", "") for c in flat.columns]
    flat.insert(0, "Station", flat.index)
    table(flat, "Hourly values removed by the flatline rule, per pollutant (whole file). "
                "Gangapur Road dominates: its NO2, NH3, O3 and CO analysers have long "
                "near-constant periods (e.g. O3 at 45.0 µg/m³ in December 2021-March 2022, "
                "NO2 at 6.28 µg/m³ in March-June 2023, NO2 at 3.85 µg/m³ from late "
                "August 2025).", font=8)
    para("Detected sensor faults that remain uncertain are reported, not hidden: in "
         "mid-2025 CO at Gangapur Road and MIDC Ambad varies only within a few 0.01 mg/m³ "
         "steps. These months pass the flatline rule and are kept.")
    h("5.2 Missing values", 2)
    figure(F["out"], "fig01_missing_values.png", "Missing hourly values per station and "
           "pollutant before imputation (all data up to the end of the 2026 horizon).", 5.8)
    mp = miss.map(lambda v: f"{v:.1f}")
    mp.insert(0, "Station", mp.index)
    table(mp, "Percentage of missing hourly values before imputation.")
    h("5.3 KNN imputation", 2)
    para("The pollutants are first standardised, then imputed. A missing value is filled "
         "with the inverse-distance-weighted mean of that variable in the k = 5 most similar "
         "hours, using the NaN-Euclidean distance over the seven standardised pollutants of "
         "the same hour [2], [5], [6].")
    bullets([
        "One StandardScaler and one KNNImputer per station and mode, fitted on that "
        "station's training rows only, then applied unchanged to validation, test and "
        "forecast inputs. Every training row is a candidate donor; no sub-sample is drawn.",
        "Many missing values occur in hours where all seven pollutants are missing (a "
        "whole-station outage). KNN has nothing to compare such an hour with, so it "
        "receives the training mean of each pollutant.",
        "Imputed values are model inputs only. Training targets are never imputed (the loss "
        "is masked) and every metric uses measured values only.",
        "Hours with every pollutant missing get the training mean; windows with more than "
        "half of their pollutant cells imputed are discarded.",
    ])
    g = imp.groupby(["station", "split"], sort=False)[["missing_before", "missing_after"]] \
        .sum().reset_index()
    g.columns = ["Station", "Split", "Missing cells before", "Missing cells after"]
    table(g, "Missing pollutant cells before and after KNN imputation, 2026-forecast mode.")


def eda_section():
    h("6. Exploratory Data Analysis")
    figure(F["out"], "fig02_historical_trends.png", "Monthly means of hourly observations per "
           "pollutant and station. Gaps are periods with no observations.", 4.9)
    para("Particulate matter follows a seasonal cycle at all stations: low during the "
         "monsoon (June-September) and high in winter. This cycle is what a year-ahead "
         "forecast can realistically capture. Gangapur Road's gas analysers (NO2, NH3, O3, "
         "CO) have long near-constant periods, removed by the flatline rule, and their "
         "levels shift between years, so its older history is a weaker guide to 2025-2026.")
    figure(F["out"], "fig03_station_comparison.png", "Distribution of hourly concentrations "
           "by station (box: interquartile range; whiskers: 1.5 IQR; outliers hidden).", 6.3)


def architecture_section():
    h("7. CNN-LSTM Architecture")
    para("Two causal one-dimensional convolutions (64 filters, kernel 3) extract local "
         "temporal motifs from the 24-hour window; max-pooling halves the sequence; an LSTM "
         "[1] with 64 units summarises it; dropout [9] regularises; two dense layers output "
         "the seven pollutants. Inputs per hour: the seven standardised, KNN-imputed "
         "pollutants plus sine and cosine encodings of hour-of-day, day-of-week and month "
         "(13 features). Calendar encodings are known for any future hour.")
    arch = pd.DataFrame([
        ["Input", "(24, 13)", "24 hours x 13 features", "0"],
        ["Conv1D", "(24, 64)", "64 filters, kernel 3, causal padding, ReLU", "2,560"],
        ["Conv1D", "(24, 64)", "64 filters, kernel 3, causal padding, ReLU", "12,352"],
        ["MaxPooling1D", "(12, 64)", "pool size 2", "0"],
        ["LSTM", "(64)", "64 units, final hidden state", "33,024"],
        ["Dropout", "(64)", "rate 0.2", "0"],
        ["Dense", "(32)", "ReLU", "2,080"],
        ["Dense (output)", "(7)", "linear, one unit per pollutant", "231"],
    ], columns=["Layer", "Output shape", "Configuration", "Parameters"])
    table(arch, "CNN-LSTM layers (50,247 trainable parameters per station model).")


def methodology_section():
    h("8. Methodology")
    h("8.1 Two modes", 2)
    m = pd.DataFrame([
        [MODES[k], r["meta"]["training_cutoff"], r["meta"]["horizon_start"],
         r["meta"]["horizon_end"], r["meta"]["horizon_hours"]] for k, r in
        (("backtest_2025", B), ("forecast_2026", F))],
        columns=["Mode", "Training cutoff", "Horizon start", "Horizon end", "Horizon hours"])
    table(m, "Run configuration. Each mode trains its own models, scalers and imputers and "
             "writes to its own output folder.", font=8)
    para("Within each mode, validation is the six months before the forecast origin and "
         "training is everything earlier: the 2025 backtest fits the network on data up to "
         "30 June 2024 and validates on July-December 2024; the 2026 run fits on data up to "
         "30 June 2025 and validates on July-December 2025. Scaler, imputer, network weights, "
         "early stopping and "
         "the choice of target design use only data up to the training cutoff. Horizon rows "
         "are used for two things only: as ground truth, and as inputs to the one-step "
         "evaluation.")
    h("8.2 Two kinds of prediction", 2)
    bullets([
        "Recursive forecast (the 2026 forecast proper). The model starts from the 24 hours "
        "before the origin, predicts the next hour, appends that prediction (with the known "
        "calendar features) to its input window, drops the oldest hour and repeats for all "
        "8,760 hours. Only the starting window is real data, so errors accumulate and, beyond "
        "the first days, the forecast can only carry the seasonal and diurnal pattern the "
        "model has learned. It is compared with three baselines that use the same "
        "information: hour-of-day by month climatology, the same hour 52 weeks earlier, and "
        "the last value before the origin.",
        "Rolling one-step evaluation (the task the network was trained for). For each "
        "measured horizon hour with a valid window, the frozen model predicts it from the "
        "measured previous 24 hours. This needs measurements up to the previous hour, so it "
        "exists only where the data exist. It is compared with persistence (the last "
        "observed value) on the same rows.",
    ])
    para("Predicted concentrations are clipped at zero, both before being fed back and on "
         "output; a pollutant concentration cannot be negative. Every horizon hour is "
         "exported. Each row records whether a measurement exists (observed), a value was "
         "rejected by the cleaning rules (removed_by_quality_check), the station file has "
         "the hour but no value (missing_value), or the file has no row for the hour "
         "(no_row_in_file). Ground truth is never imputed or interpolated, and only "
         "observed hours are scored.")
    h("8.3 Training and target design", 2)
    cfg = pd.DataFrame([
        ["Loss", "Masked MSE on standardised targets (missing targets excluded)"],
        ["Optimiser", "Adam [8], learning rate 0.001"],
        ["Batch size / maximum epochs", f"{F['meta']['batch_size']} / {F['meta']['max_epochs']}"],
        ["Early stopping", f"validation loss, patience {F['meta']['patience']}, best weights"],
        ["Random seed", f"{F['meta']['seed']} (Python, NumPy, TensorFlow)"],
        ["Software", f"Python {V['python']}, TensorFlow {V['tensorflow']}, scikit-learn "
                     f"{V['scikit-learn']}, pandas {V['pandas']}, NumPy {V['numpy']}"],
    ], columns=["Setting", "Value"])
    table(cfg, "Training configuration (same for all stations and modes; no hyperparameter "
               "search).", font=8)
    para("Two parameterisations of the same network are trained per station: level "
         "(predict the standardised concentration) and change (predict the difference from "
         "the last hour of the input window, then add it back). Training, one-step and "
         "recursive prediction all use that same base value; in the recursive forecast it "
         "is the previous prediction. The design for the recursive forecast is chosen by a "
         "recursive forecast of the validation half-year. The design for the one-step "
         "evaluation is chosen by one-step validation loss. Neither choice sees the horizon.")
    for key, r in (("backtest_2025", B), ("forecast_2026", F)):
        s = r["sel"].reset_index()[["station", "val_loss_persistence", "val_loss_level",
                                    "val_loss_change", "val_recursive_mse_level",
                                    "val_recursive_mse_change", "selected_design",
                                    "one_step_design"]]
        s.columns = ["Station", "1-step val: persistence", "1-step val: level",
                     "1-step val: change", "Recursive val: level", "Recursive val: change",
                     "Recursive design", "One-step design"]
        table(s, f"{MODES[key]}: model selection on validation data (MSE, standardised units).",
              font=7)
    figure(F["out"], "fig04_loss_curves.png", "2026-forecast mode: training and validation "
           "loss of the recursively selected design per station.", 6.5)


def split_section():
    h("9. Chronological Splits")
    para("Random splitting would score the model on near-copies of its training hours and "
         "let it learn from the period it is tested on [10]. Every split here is "
         "chronological, and windows are assigned by the timestamp of the predicted hour.")
    sp = pd.concat([B["splits"].assign(mode="2025 backtest"),
                    F["splits"].assign(mode="2026 forecast")])[
        ["mode", "station", "split", "start", "end", "hourly_rows", "pm25_observed"]]
    sp.columns = ["Mode", "Station", "Split", "Start", "End", "Hourly rows", "PM2.5 observed"]
    sp["Start"] = sp["Start"].fillna("-")
    sp["End"] = sp["End"].fillna("-")
    table(sp, "Splits per mode and station. Test = the horizon rows present in the file.",
          font=7)


def metrics_section():
    h("10. Evaluation Metrics")
    para("For n measured hours with values yᵢ, predictions ŷᵢ and measured mean ȳ:")
    equation("MSE = (1/n) Σ (yᵢ − ŷᵢ)²")
    equation("RMSE = √MSE")
    equation("MAE = (1/n) Σ |yᵢ − ŷᵢ|")
    equation("R² = 1 − Σ (yᵢ − ŷᵢ)² / Σ (yᵢ − ȳ)²")
    bullets([
        "All metrics are in original units, after inverse scaling, on measured hours only. "
        "Every method is scored on exactly the same hours.",
        "n_valid is the number of measured hours scored. Coverage is n_valid as a share of "
        "the 8,760 horizon hours.",
        "R² is reported as undefined with fewer than two hours or zero variance, and "
        "flagged unreliable with fewer than 168 hours (one week).",
        "Recursive skill = 1 − RMSE_CNN-LSTM / RMSE_climatology; one-step skill = "
        "1 − RMSE_CNN-LSTM / RMSE_persistence. Positive means the CNN-LSTM is better.",
        "Hours removed by the cleaning rules are never scored, for any method.",
        "As a second safeguard, a series whose remaining measured horizon values take fewer "
        "than 100 distinct values is quality-flagged. It is reported but left out of the "
        "pooled per-pollutant summary and the win counts. Pollutants are never pooled "
        "together.",
        "R² is never read on its own: hourly pollution is persistent, so even the "
        "persistence baseline reaches a high R². The skill against a baseline is the "
        "number that shows whether the model adds anything.",
    ])


def recursive_tables(r, label):
    t = r["rec"]
    x = t[["station", "pollutant", "n_valid", "coverage_pct", "CNN-LSTM RMSE",
           "CNN-LSTM MAE", "CNN-LSTM R2", "Climatology RMSE", "Same hour last year RMSE",
           "Persistence RMSE", "RMSE skill vs climatology", "r2_flag"]].copy()
    x["pollutant"] = flag(t)
    x.columns = ["Station", "Pollutant", "n", "Cov. %", "RMSE", "MAE", "R²", "RMSE clim.",
                 "RMSE last yr", "RMSE persist.", "Skill vs clim.", "R² flag"]
    table(x, f"{label}: recursive forecast per station and pollutant, measured hours only "
             "(µg/m³; CO in mg/m³). * = quality-flagged series; '-' = no measured hours.",
          font=6)
    o = r["rec_all"][["pollutant", "n_valid", "CNN-LSTM MSE", "CNN-LSTM RMSE", "CNN-LSTM MAE",
                      "CNN-LSTM R2", "Climatology RMSE", "RMSE skill vs climatology",
                      "stations_included"]].copy()
    o.columns = ["Pollutant", "n", "MSE", "RMSE", "MAE", "R²", "RMSE clim.", "Skill vs clim.",
                 "Stations"]
    table(o, f"{label}: recursive forecast, unflagged stations pooled per pollutant.", font=7)


def one_step_tables(r, label):
    t = r["one"]
    x = t[["station", "pollutant", "n_valid", "CNN-LSTM MSE", "CNN-LSTM RMSE", "CNN-LSTM MAE",
           "CNN-LSTM R2", "Persistence RMSE", "RMSE skill vs persistence", "r2_flag"]].copy()
    x["pollutant"] = flag(t)
    x.columns = ["Station", "Pollutant", "n", "MSE", "RMSE", "MAE", "R²", "RMSE persist.",
                 "Skill vs persist.", "R² flag"]
    table(x, f"{label}: rolling one-step predictions per station and pollutant. "
             "* = quality-flagged series; '-' = no measured hours.", font=6)
    o = r["one_all"][["pollutant", "n_valid", "CNN-LSTM RMSE", "CNN-LSTM R2",
                      "Persistence RMSE", "RMSE skill vs persistence", "stations_included"]]
    o = o.copy()
    o.columns = ["Pollutant", "n", "RMSE", "R²", "RMSE persist.", "Skill vs persist.",
                 "Stations"]
    table(o, f"{label}: one-step predictions, unflagged stations pooled per pollutant.", font=7)


def backtest_section():
    h("11. Results: 2025 Backtest")
    para("Models trained on data up to 31 December 2024 forecast every hour of 2025, which "
         "is measured at all four stations. This is the fair test of how the procedure "
         "behaves on a year it has never seen.")
    h("11.1 Recursive year-ahead forecast", 2)
    recursive_tables(B, "2025")
    n_pos = int((B_rec_skill > 0).sum())
    para(f"Pooled RMSE skill against climatology ranges from {B_rec_skill_txt}: the "
         f"CNN-LSTM beats climatology for {n_pos} of {len(POLL)} pollutants. A year-ahead "
         "recursive forecast cannot know the weather of a given day, so after the first "
         "hours its value lies in reproducing the seasonal and daily cycle, and climatology "
         "is the baseline it has to match. It does not: the CNN-LSTM beats climatology for "
         f"only {B_rec_wins} of {B_rec_n} station-pollutant series. Feeding predictions back "
         "lets the network settle on its own level, which can be far from the measurements "
         f"for months. The worst series is {worst_series(B)}.")
    for s in STATIONS:
        figure(B["out"], f"fig05_forecast_{s.replace(' ', '_')}.png",
               f"{s}: 2025 recursive forecast (daily means) against measurements.", 6.0)
    figure(B["out"], "fig05b_first_two_weeks_pm25.png", "PM2.5 in the first 14 days of "
           "2025: the recursive forecast drifts from the measurements within days.", 6.0)
    figure(B["out"], "fig09_rmse_by_month.png", "Recursive forecast RMSE by month of 2025 "
           "(forecast lead), unflagged stations pooled.", 6.5)
    figure(B["out"], "fig08_station_metrics.png", "2025 recursive RMSE per station and "
           "baseline, and RMSE skill against climatology.", 6.5)
    h("11.2 Rolling one-step evaluation", 2)
    one_step_tables(B, "2025")
    para(f"With the measured previous 24 hours as input, pooled RMSE skill against "
         f"persistence ranges from {B_one_skill_txt}. Persistence is a strong baseline one "
         "hour ahead because consecutive hourly averages are highly correlated.")
    figure(B["out"], "fig10_scatter_one_step.png", "2025 one-step predictions against "
           "measurements. Dashed: perfect prediction.", 5.4)


def forecast_section():
    h("12. Results: 2026 (Partial-Year Evaluation)")
    para("Models trained on data up to 31 December 2025 forecast every hour of 2026 at all "
         "four stations (8,760 hours each, 1 January 00:00 to 31 December 23:00). The forecast "
         "file contains every hour. Only hours labelled observed are scored: January-June "
         "and August at MIDC Ambad, Gangapur Road and Hirwadi. July and September-December "
         "2026, and all of 2026 at Pandav Nagari, are forecasts without ground truth. "
         "Because those months are missing, the 2026 scores do not cover a full seasonal "
         "cycle (the monsoon months July-September are mostly absent).")
    h("12.1 Plausibility of the unmeasured months", 2)
    para("Metrics cannot judge the months without measurements, so Table "
         f"{TAB_N[0] + 1} compares the monthly mean PM2.5 forecast with the station's "
         "climatology for every month of 2026. Large ratios in either direction mark drift of "
         "the recursive forecast rather than an expected change in air quality.")
    t = pd.concat({"Forecast": pm_monthly["cnn_lstm_pred"],
                   "Climatology": pm_monthly["climatology_pred"]}).round(0)
    t = t.swaplevel().sort_index(level=0, sort_remaining=False).reindex(
        [(st, k) for st in STATIONS for k in ("Forecast", "Climatology")])
    t.columns = [pd.Timestamp(2026, m, 1).strftime("%b") for m in t.columns]
    t.insert(0, "Series", [k for _, k in t.index])
    t.insert(0, "Station", [st for st, _ in t.index])
    table(t.reset_index(drop=True), "2026 monthly mean PM2.5 (µg/m³): recursive CNN-LSTM "
          "forecast against climatology.", font=7)
    off = pm_ratio.stack()
    off = off[(off > 2) | (off < 0.5)]
    months = {st: ", ".join(pd.Timestamp(2026, m, 1).strftime("%b") for m in g.index
                            .get_level_values(1)) for st, g in off.groupby(level=0)}
    designs = F["sel"]["selected_design"]
    top = (off.apply(lambda v: max(v, 1 / v))).idxmax()
    para("Months where the forecast is more than twice, or less than half of, "
         "climatology: " + "; ".join(f"{st} ({ms}; {designs[st]} design)"
                                     for st, ms in months.items())
         + f". The most extreme is {top[0]} in {pd.Timestamp(2026, top[1], 1):%B}: "
         f"{pm_monthly['cnn_lstm_pred'].loc[top]:.0f} µg/m³ against a climatology of "
         f"{pm_monthly['climatology_pred'].loc[top]:.0f}. Both target designs drift: each "
         "prediction becomes an input for the next, so small biases accumulate over "
         "thousands of steps and carry the forecast away from realistic levels. These values, "
         "and the AQI derived from them, should not be used as an estimate of 2026 air "
         "quality; the climatology column of the same file is the better year-ahead "
         "estimate.")
    h("12.2 Recursive forecast on measured hours", 2)
    recursive_tables(F, "2026")
    for s in STATIONS:
        figure(F["out"], f"fig05_forecast_{s.replace(' ', '_')}.png",
               f"{s}: 2026 recursive forecast (daily means), with measurements where they "
               "exist. Grey: no measurement.", 6.0)
    figure(F["out"], "fig06_scatter_recursive.png", "2026 recursive forecast against "
           "measured hourly values.", 5.4)
    h("12.3 Rolling one-step evaluation on measured hours", 2)
    one_step_tables(F, "2026")
    figure(F["out"], "fig10_scatter_one_step.png", "2026 one-step predictions against "
           "measurements. Pandav Nagari has no 2026 measurements.", 5.4)


def aqi_section():
    h("13. Predicted AQI")
    para("The model predicts pollutant concentrations; AQI is derived from them afterwards "
         "with the CPCB National AQI method [11]. Sub-indices use 24-hour means for PM2.5, "
         "PM10, NO2, SO2 and NH3 and the maximum 8-hour mean for CO and O3. The AQI is the "
         "largest sub-index and is computed only when at least three sub-indices, including "
         "PM2.5 or PM10, are available. A measured daily mean needs at least 16 measured "
         "hours. The Severe band has no published upper concentration, so it is extended "
         "linearly with the width of the Very Poor band and capped at 500. Predicted AQI is a "
         "model estimate from the recursive forecast. Measured AQI is computed from the "
         "station data with the same rules, not taken from an official bulletin.")
    for key, r in (("backtest_2025", B), ("forecast_2026", F)):
        a = r["aqi"].copy()
        a.columns = ["Station", "Forecast days", "Days with measured AQI",
                     "Category agreement %", "AQI RMSE", "Mean predicted AQI",
                     "Mean measured AQI"]
        table(a, f"{MODES[key]}: daily AQI from the recursive forecast against measured AQI.",
              font=7)
        figure(r["out"], "fig11_daily_aqi.png", f"{MODES[key]}: daily AQI derived from the "
               "forecast, with measured AQI on valid days. Bands: CPCB categories.", 6.3)


def discussion_section():
    h("14. Discussion")
    para("Leakage checks. A separate script (scripts/verify_outputs.py) runs for each mode "
         "and re-derives the results from the raw CSV files and the saved models. It "
         "replays the first 48 hours of every station's recursive forecast from raw data "
         "truncated at the cutoff and reproduces the export, so the forecast does not "
         "depend on any later measurement. It rebuilds 20 random one-step input windows per "
         "station from the raw files, confirming that each prediction uses exactly the 24 "
         "preceding hours. It also checks that the scaler means equal the training-period "
         "means, that every horizon hour is present once, that missing hours carry no "
         "actual value, and that every metric matches a recomputation from the exported "
         "predictions.")
    for key, r in (("backtest_2025", B), ("forecast_2026", F)):
        flagged = r["rec"][r["rec"].quality_flag]
        names = ", ".join(f"{s} {p}" for s, p in zip(flagged.station, flagged.pollutant))
        para(f"Quality-flagged series, {MODES[key]}: "
             + (f"{names} still take fewer than 100 distinct measured values after cleaning. "
                "They are reported in every table but excluded from pooled summaries and win "
                "counts." if names else "none after cleaning."))
    q = pd.concat([r["rec"].assign(mode=MODES[k]) for k, r in
                   (("backtest_2025", B), ("forecast_2026", F))])
    q = q.pivot_table(index=["mode", "station"], columns="pollutant",
                      values="hours_removed_by_quality_check", aggfunc="sum")[POLL]
    q = q.reset_index()
    q.columns = ["Mode", "Station", *POLL]
    table(q, "Horizon hours whose measurement was removed by the cleaning rules and is "
             "therefore not scored (out of 8,760 per series).", font=8)
    para("Recursive versus one-step. The two scores answer different questions. The "
         "one-step score says how well the network corrects persistence when fresh "
         "measurements arrive every hour. The recursive score says what can be said about a "
         "given hour months in advance, which is mostly its seasonal and diurnal expectation. "
         "The 2026 forecast file is the recursive forecast. The one-step results show what an "
         "operational next-hour system would achieve once 2026 data arrives.")


def limitations_section():
    h("15. Limitations")
    bullets([
        "2026 ground truth covers January-June and August at three stations only; July and "
        "September-December are unmeasured, and Pandav Nagari has no 2026 data. The 2026 "
        "scores therefore miss most of the monsoon and the post-monsoon season.",
        "A recursive hourly forecast months ahead carries no information about specific "
        "weather events. Its hourly values should be read as an expected pattern, not as a "
        "prediction of individual hours.",
        "No uncertainty intervals: one training run per model (seed 42). Recursive errors "
        "accumulate with lead time, as the RMSE-by-month figure shows.",
        "Three stations have about two and a half years of training data (mid-2023 to the "
        "cutoff), and Gangapur Road's long history includes years whose sensor behaviour "
        "differs from 2026.",
        "The model uses no weather inputs (temperature, humidity, wind, rain), although "
        "they are measured at three of the four stations. Weather drives much of the "
        "day-to-day change in pollution, and a year-ahead forecast could only use it with "
        "weather forecasts, which this project does not have.",
        "The cleaning rules (12-hour stuck values, CO zeros, flatline days) and the "
        "distinct-value flag are heuristics; some degraded readings may remain, such as "
        "low-resolution CO in mid-2025.",
        "Derived AQI depends on the CPCB breakpoints as implemented here, and on the "
        "assumption for the open-ended Severe band.",
    ])


def conclusion_section():
    h("16. Conclusion")
    para("The pipeline produces a leakage-free, timestamp-aligned hourly forecast of all of "
         "2026 for all four stations, and scores it only where measurements exist. The "
         "independent verifier and the test suite confirm the horizon, the labelling of "
         "missing ground truth, the absence of post-cutoff information, and every metric.")
    para("Accuracy is a separate question, and the answer is not flattering. As a year-"
         "ahead forecast, the recursive CNN-LSTM is worse than the month-by-hour climatology "
         f"for {B_rec_worse_txt}, pooled, in the 2025 backtest ({B_rec_skill_txt}) and for "
         f"{F_rec_worse_txt} in the measured part of 2026 ({F_rec_skill_txt}). Per series it "
         f"beats climatology in only {B_rec_wins} of {B_rec_n} (2025) and {F_rec_wins} of "
         f"{F_rec_n} (2026) cases. For planning-horizon estimates of 2026, "
         "the climatology column of the same forecast file is the better number. The network "
         "earns its place one hour ahead, modestly: with the previous 24 measured hours as "
         f"input, it beats persistence for {B_one_wins} of {B_one_n} series in 2025 and "
         f"{F_one_wins} of {F_one_n} in 2026 (median RMSE reduction {100 * B_one_median:.1f}% "
         f"and {100 * F_one_median:.1f}%). Pooled per pollutant, its RMSE skill against "
         f"persistence is {B_one_skill_txt} in 2025 and {F_one_skill_txt} in 2026.")
    para("Future work: train a direct multi-horizon model, or blend the network with "
         "climatology using weights chosen on validation, rather than feeding a one-step "
         "model its own output for 8,760 steps; add quality-controlled meteorological "
         "inputs; report variance over several seeds; and re-score 2026 when July and "
         "September-December data and Pandav Nagari's 2026 data become available.")


def reproduce_section():
    h("17. Reproducing the Results")
    for line in [
        "NASHIK_MODE=backtest_2025 python nashik_aq_pipeline.py",
        "NASHIK_MODE=forecast_2026 python nashik_aq_pipeline.py",
        "python scripts/verify_outputs.py backtest_2025",
        "python scripts/verify_outputs.py forecast_2026",
        "python -m unittest discover tests -v",
        "python scripts/build_report.py",
    ]:
        p = doc.add_paragraph(line)
        p.runs[0].font.name = "Consolas"
        p.runs[0].font.size = Pt(9)
    para("Outputs are written to outputs/backtest_2025/ and outputs/forecast_2026/. "
         "run_metadata.json in each folder records the cutoff, horizon, features, target "
         "order, selected designs, model files, evaluation coverage, software versions and "
         f"git commit ({F['meta']['git_commit']}).")


def references_section():
    h("18. References")
    refs = [
        "S. Hochreiter and J. Schmidhuber, \"Long Short-Term Memory,\" Neural Computation, "
        "vol. 9, no. 8, pp. 1735-1780, 1997, doi: 10.1162/neco.1997.9.8.1735.",
        "O. Troyanskaya et al., \"Missing value estimation methods for DNA microarrays,\" "
        "Bioinformatics, vol. 17, no. 6, pp. 520-525, 2001, "
        "doi: 10.1093/bioinformatics/17.6.520.",
        "C.-J. Huang and P.-H. Kuo, \"A Deep CNN-LSTM Model for Particulate Matter (PM2.5) "
        "Forecasting in Smart Cities,\" Sensors, vol. 18, no. 7, art. 2220, 2018, "
        "doi: 10.3390/s18072220.",
        "D. Qin et al., \"A Novel Combined Prediction Scheme Based on CNN and LSTM for Urban "
        "PM2.5 Concentration,\" IEEE Access, vol. 7, pp. 20050-20059, 2019, "
        "doi: 10.1109/ACCESS.2019.2897028.",
        "H. Junninen et al., \"Methods for imputation of missing values in air quality data "
        "sets,\" Atmospheric Environment, vol. 38, no. 18, pp. 2895-2907, 2004, "
        "doi: 10.1016/j.atmosenv.2004.02.026.",
        "F. Pedregosa et al., \"Scikit-learn: Machine Learning in Python,\" Journal of "
        "Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
        "Y. LeCun, L. Bottou, Y. Bengio and P. Haffner, \"Gradient-based learning applied to "
        "document recognition,\" Proceedings of the IEEE, vol. 86, no. 11, pp. 2278-2324, "
        "1998, doi: 10.1109/5.726791.",
        "D. P. Kingma and J. Ba, \"Adam: A Method for Stochastic Optimization,\" "
        "arXiv:1412.6980, 2014 (ICLR 2015).",
        "N. Srivastava et al., \"Dropout: A Simple Way to Prevent Neural Networks from "
        "Overfitting,\" Journal of Machine Learning Research, vol. 15, pp. 1929-1958, 2014.",
        "C. Bergmeir and J. M. Benítez, \"On the use of cross-validation for time series "
        "predictor evaluation,\" Information Sciences, vol. 191, pp. 192-213, 2012, "
        "doi: 10.1016/j.ins.2011.12.028.",
        "Central Pollution Control Board, \"National Air Quality Index,\" Ministry of "
        "Environment, Forest and Climate Change, Government of India, 2014.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph(f"[{i}] {ref}")
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)


for section in (front_matter, intro_sections, dataset_section, preprocessing_section,
                eda_section, architecture_section, methodology_section, split_section,
                metrics_section, backtest_section, forecast_section, aqi_section,
                discussion_section, limitations_section, conclusion_section,
                reproduce_section,
                references_section):
    section()
doc.save(REPORT)
print(f"{REPORT}: {FIG_N[0]} figures, {TAB_N[0]} tables")
